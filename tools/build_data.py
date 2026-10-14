#!/usr/bin/env python3
"""Build MaklonGo data files.

Inputs (folder `source/`, NOT committed – contains cost prices):
  source/data_produk.xlsx       Data Produk dan Deskripsi (Shopee export)
  source/discount_nominate.xlsx harga produk (Harga Awal / Harga diskon)
  source/master_bahan.csv       master bahan baku + harga
  source/kemasan_drive.json     daftar file gambar kemasan dari Google Drive
                                [[drive_id, file_name, mime], ...]
  source/thumbs/<drive_id>.png  (opsional) gambar kemasan yang sudah diunduh

Curated, public inputs (committed, tanpa harga):
  data/bases.json    daftar base produk (nama baru, bahan aktif bawaan, varian)
  data/actives.json  daftar bahan aktif + kategori + kode bahan di master

Outputs:
  assets/js/catalog.js       katalog publik (TANPA harga) untuk web customer
  data/prices.enc.json       harga terenkripsi (AES-GCM, kunci = password admin)
  assets/img/kemasan/*.webp  gambar kemasan terkompresi
  source/prices.plain.json   salinan harga tidak terenkripsi (lokal saja)

Usage:
  MAKLONGO_ADMIN_PASSWORD='...' python3 tools/build_data.py [--images]
"""
import base64
import csv
import json
import os
import re
import sys
import time
from pathlib import Path

import openpyxl
from cryptography.hazmat.primitives import hashes
from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "source"
DATA = ROOT / "data"
IMG_DIR = ROOT / "assets" / "img" / "kemasan"

PBKDF2_ITER = 250_000

DEFAULT_SETTINGS = {
    "moq": 50,
    "activeDosePct": 1,        # asumsi pemakaian bahan aktif 1% w/v
    "activeScrapFactor": 1.35, # +35% scrap factor produksi
    "packagingMarkup": 1.35,   # harga kemasan di nama file x 135%
    "bpomLabPerSku": 3_500_000,
    "halalPerSku": 3_500_000,
    "stickerPerPcs": 0,        # isi di dashboard admin
}


# ---------------------------------------------------------------- base produk
def load_base_prices():
    wb = openpyxl.load_workbook(SRC / "discount_nominate.xlsx", data_only=True)
    ws = wb.active
    prices = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        code = str(r[0]).strip()
        prices.setdefault(code, []).append({
            "variasi": (r[4] or "").strip(),
            "hargaAwal": float(r[6] or 0),
            "hargaDiskon": float(r[7] or 0),
        })
    return prices


def build_bases():
    bases = json.loads((DATA / "bases.json").read_text())
    prices = load_base_prices()
    public, secret = [], {}
    for b in bases:
        rows = prices.get(b["shopee"])
        if not rows:
            sys.exit(f"Harga tidak ditemukan untuk base {b['id']} ({b['shopee']})")
        # Produk dengan beberapa ukuran (250ml / 1L): pakai varian yang volumenya sama dengan base
        row = next((x for x in rows if re.search(rf"\b{b['vol']}\s*ml\b", x["variasi"], re.I)), None)
        if row is None:
            row = max(rows, key=lambda x: x["hargaDiskon"])
        secret[b["id"]] = {
            "price": round(row["hargaDiskon"]),
            "priceAwal": round(row["hargaAwal"]),
            "vol": b["vol"],
            "unit": b["unit"],
            "shopee": b["shopee"],
        }
        public.append({k: v for k, v in b.items() if k != "shopee"})
    return public, secret


# --------------------------------------------------------------- bahan aktif
def build_actives():
    spec = json.loads((DATA / "actives.json").read_text())
    master = {}
    with open(SRC / "master_bahan.csv", newline="", encoding="utf-8-sig") as f:
        for r in csv.DictReader(f, delimiter=";"):
            master[r["kode"]] = r
    public, secret = [], {}
    for a in spec["items"]:
        cands = []
        for code in a["bb"]:
            r = master.get(code)
            if not r:
                sys.exit(f"Kode bahan {code} untuk {a['name']} tidak ada di master")
            per_unit = float(r["harga"]) / float(r["per_qty"]) / float(r["densitas"] or 1)
            cands.append({
                "kode": code, "nama": r["nama"], "supplier": r["supplier"],
                "harga": float(r["harga"]), "perQty": float(r["per_qty"]),
                "satuan": r["satuan"], "pricePerGram": round(per_unit, 4),
            })
        # Harga tertinggi dipakai sebagai default (konservatif); bisa diubah di admin
        chosen = max(cands, key=lambda c: c["pricePerGram"])
        secret[a["id"]] = {"pricePerGram": chosen["pricePerGram"], "source": chosen, "alternatives": cands}
        public.append({"id": a["id"], "name": a["name"], "cats": a["cats"]})
    return spec["categories"], public, secret


# ------------------------------------------------------------------ kemasan
ACRONYMS = {"PET", "PE", "PP", "HDPE", "RF", "RFS", "SBR", "LDPE", "BPA", "UV", "II"}
COLORS = ["PUTIH", "HITAM", "CLEAR", "BENING", "AMBER", "COKLAT", "COKELAT", "PINK", "BIRU", "HIJAU",
          "MERAH", "MARON", "MAROON", "KUNING", "FROSTED", "GOLD", "SILVER", "NATURAL", "TRANSPARAN",
          "UNGU", "ABU", "TOSCA", "ORANGE", "CREAM", "NAVY", "ROSE GOLD", "TOSKA", "MAGENTA", "NUDE"]
VOL_RE = re.compile(r"^(\d+(?:[.,]\d+)?)\s*(ml|l|lt|liter|g|gr|gram)$", re.I)
PRICE_SEG_RE = re.compile(r"^(?:R[pPo]?\s*[.,]?)?\d[\d.,]*$")


def titleize(s):
    out = []
    for w in s.split():
        out.append(w if w.upper() in ACRONYMS else w.capitalize())
    return " ".join(out)


def pack_type(desc, has_vol):
    d = desc.upper()
    if not has_vol:
        return "aksesoris"
    if re.search(r"\bPOT\b|\bJAR\b|CUSHION", d):
        return "pot"
    if re.search(r"\bTUBE\b", d):
        return "tube"
    if re.search(r"SPRAY|TRIGER|TRIGGER|KISPRAY|SEMPROT|MIST|NOZZ?EL", d):
        return "spray"
    if re.search(r"PUMP|POMPA", d):
        return "pump"
    if re.search(r"PIPET|DROPPER|TETES", d):
        return "pipet"
    if re.search(r"FLIP", d):
        return "fliptop"
    if re.search(r"JERIGEN|GALON", d):
        return "jerigen"
    if re.search(r"ROLL", d):
        return "rollon"
    return "botol"


def parse_pack(name):
    s = re.sub(r"\.(png|jpe?g|webp)$", "", name.strip(), flags=re.I).strip()
    m = re.match(r"(W\d+)\s*-?\s*", s)
    if not m:
        return None
    code = m.group(1)
    rest = s[m.end():]
    note = " ".join(x.strip() for x in re.findall(r"\[([^\]]*)\]", rest))
    rest = re.sub(r"\s*\[[^\]]*\]", "", rest).strip()
    segs = [x.strip() for x in re.split(r"\s*-\s*", rest) if x.strip()]
    price = None
    num = lambda x: int(re.sub(r"\D", "", x))
    if len(segs) >= 2 and PRICE_SEG_RE.match(segs[-1]):
        nums = [num(segs.pop())]
        if len(segs) >= 2 and re.match(r"^R[pPo]?\s*\d[\d.,]*$", segs[-1]):
            nums.append(num(segs.pop()))  # rentang harga "Rp1,800 - 1,900"
        price = max(nums)
    elif segs:
        m2 = re.search(r"\s+R[pPo]\s*(\d[\d.,]*)$", segs[-1])  # "... SPRAY Rp3,200"
        if m2:
            price = num(m2.group(1))
            segs[-1] = segs[-1][:m2.start()].strip()
    vol = unit = None
    if segs:
        vm = VOL_RE.match(segs[0].replace(" ", ""))
        if vm:
            v = float(vm.group(1).replace(",", "."))
            u = vm.group(2).lower()
            if u in ("l", "lt", "liter"):
                v, u = v * 1000, "ml"
            elif u in ("g", "gr", "gram"):
                u = "g"
            vol, unit = (int(v) if v.is_integer() else v), u
            segs = segs[1:]
    desc = " - ".join(segs)
    up = desc.upper()
    colors = [c for c in COLORS if re.search(rf"\b{c}\b", up)]
    if vol is None and len(segs) >= 2:
        name_, variasi = segs[0], " - ".join(segs[1:])
    else:
        name_, variasi = desc, ", ".join(titleize(c) for c in colors) or "-"
    if note:
        variasi = f"{variasi} ({note[0].upper() + note[1:]})"
    return {
        "code": code, "name": titleize(name_), "variasi": titleize(variasi) if vol is None else variasi,
        "vol": vol, "unit": unit, "type": pack_type(desc, vol is not None), "filePrice": price,
    }


def build_packaging(settings, make_images):
    files = json.loads((SRC / "kemasan_drive.json").read_text())
    items, secret, seen = [], {}, {}
    if make_images:
        from PIL import Image
        IMG_DIR.mkdir(parents=True, exist_ok=True)
    for drive_id, fname, _mime in files:
        p = parse_pack(fname)
        if not p:
            print("skip", fname)
            continue
        n = seen.get(p["code"], 0)
        seen[p["code"]] = n + 1
        key = p["code"] if n == 0 else f"{p['code']}-{n + 1}"
        img = f"assets/img/kemasan/{key.lower()}.webp"
        if make_images:
            src = SRC / "thumbs" / f"{drive_id}.png"
            dst = ROOT / img
            if src.exists() and not dst.exists():
                im = Image.open(src).convert("RGB")
                im.thumbnail((320, 320))
                im.save(dst, "WEBP", quality=70, method=6)
        items.append({"id": key, "code": p["code"], "name": p["name"], "variasi": p["variasi"],
                      "vol": p["vol"], "unit": p["unit"], "type": p["type"], "img": img})
        fp = p["filePrice"]
        secret[key] = {"filePrice": fp, "price": round(fp * settings["packagingMarkup"]) if fp else None,
                       "file": fname, "driveId": drive_id}
    items.sort(key=lambda x: (x["type"] == "aksesoris", x["vol"] or 0, x["code"]))
    return items, secret


# Katalog kemasan baru (MP0001…): harga supplier dari Excel, foto dari folder Drive
def mp_type(name):
    n = name.upper()
    if re.search(r"LIP|MASKARA|EYELINER|KUTEK|LIPSTIK", n):
        return "makeup"
    if re.search(r"POT|JAR|CUSHION", n):
        return "pot"
    if re.search(r"TUBE", n):
        return "tube"
    if re.search(r"ROLL ON", n):
        return "rollon"
    if re.search(r"PIPET", n):
        return "pipet"
    if re.search(r"SPRAY|PARFUM|MIST", n):
        return "spray"
    if re.search(r"PUMP|AIRLESS|FOAM", n):
        return "pump"
    if re.search(r"FLIPTOP", n):
        return "fliptop"
    return "botol"


def parse_mp_volume(v):
    # "8-10 ml" → isi 8 (batas bawah, aman), label tetap "8-10"; "3,5 gr" → 3.5 g
    m = re.match(r"^\s*([\d.,]+)(?:\s*-\s*([\d.,]+))?\s*(ml|gr|g)\s*$", str(v or ""), re.I)
    if not m:
        return None, None, None
    lo = float(m.group(1).replace(",", "."))
    unit = "ml" if m.group(3).lower() == "ml" else "g"
    label = f"{m.group(1)}-{m.group(2)}" if m.group(2) else None
    return (int(lo) if lo.is_integer() else lo), unit, label


def build_packaging_mp(settings, make_images):
    path = SRC / "kemasan_mp.xlsx"
    if not path.exists():
        return [], {}
    photos = {t.split(" - ")[0]: i for i, t in json.loads((SRC / "kemasan_mp_drive.json").read_text())}
    ws = openpyxl.load_workbook(path, data_only=True)["Produk"]
    items, secret = [], {}
    if make_images:
        from PIL import Image
        IMG_DIR.mkdir(parents=True, exist_ok=True)
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[1]:
            continue
        parent, sku, name, vol_raw = str(r[0]).strip(), str(r[1]).strip(), str(r[2]).strip(), r[3]
        vol, unit, vlabel = parse_mp_volume(vol_raw)
        opts = [str(x).strip() for x in (r[10], r[12], r[14]) if x and str(x).strip()]
        variasi = ", ".join(o if o.isupper() is False else o.title() for o in opts) or "-"
        img = f"assets/img/kemasan/{parent.lower()}.webp"
        drive_id = photos.get(parent)
        if make_images and drive_id:
            src = SRC / "thumbs" / f"{drive_id}.png"
            dst = ROOT / img
            if src.exists() and not dst.exists():
                im = Image.open(src).convert("RGB")
                im.thumbnail((320, 320))
                im.save(dst, "WEBP", quality=72, method=6)
        item = {"id": sku, "code": sku, "name": name, "variasi": variasi, "vol": vol, "unit": unit,
                "type": mp_type(name), "img": img if drive_id else None, "new": True}
        if vlabel:
            item["volLabel"] = vlabel
        items.append(item)
        harga, diskon = float(r[15] or 0), float(r[16] or 0)
        # Harga modal = kolom "Diskon" (harga diskon supplier); kolom "Harga" hanya referensi
        modal = diskon or harga
        secret[sku] = {"filePrice": round(modal), "supplierListPrice": round(harga),
                       "price": round(modal * settings["packagingMarkup"]) if modal else None,
                       "file": f"{sku} - {name} - {vol_raw}", "driveId": drive_id, "source": "Produk_MaklonGo_-_Gabungan_SKU.xlsx"}
    return items, secret


# --------------------------------------------------------------- enkripsi
def encrypt(obj, password):
    salt, iv = os.urandom(16), os.urandom(12)
    kdf = PBKDF2HMAC(algorithm=hashes.SHA256(), length=32, salt=salt, iterations=PBKDF2_ITER)
    key = kdf.derive(password.encode())
    ct = AESGCM(key).encrypt(iv, json.dumps(obj, ensure_ascii=False).encode(), None)
    b64 = lambda b: base64.b64encode(b).decode()
    return {"v": 1, "kdf": "PBKDF2-SHA256", "iter": PBKDF2_ITER, "salt": b64(salt), "iv": b64(iv), "ct": b64(ct)}


def main():
    password = os.environ.get("MAKLONGO_ADMIN_PASSWORD")
    if not password:
        sys.exit("Set MAKLONGO_ADMIN_PASSWORD")
    make_images = "--images" in sys.argv
    settings = dict(DEFAULT_SETTINGS)

    bases_pub, bases_sec = build_bases()
    cats, actives_pub, actives_sec = build_actives()
    pack_pub, pack_sec = build_packaging(settings, make_images)
    mp_pub, mp_sec = build_packaging_mp(settings, make_images)
    pack_pub, pack_sec = mp_pub + pack_pub, {**mp_sec, **pack_sec}

    stamp = time.strftime("%Y-%m-%d")
    catalog = {
        "version": stamp,
        "baseCategories": [
            {"id": "face", "name": "Perawatan Wajah", "icon": "✨"},
            {"id": "body", "name": "Perawatan Tubuh", "icon": "🧴"},
            {"id": "hair", "name": "Perawatan Rambut", "icon": "💇"},
            {"id": "hygiene", "name": "Kewanitaan", "icon": "🌸"},
            {"id": "lip", "name": "Bibir & Gigi", "icon": "💋"},
        ],
        "bases": bases_pub,
        "activeCategories": cats,
        "actives": actives_pub,
        "packagingTypes": [
            {"id": "pump", "name": "Botol Pump"},
            {"id": "spray", "name": "Botol Spray"},
            {"id": "botol", "name": "Botol Tutup"},
            {"id": "fliptop", "name": "Flip Top"},
            {"id": "pipet", "name": "Botol Pipet"},
            {"id": "pot", "name": "Pot / Jar"},
            {"id": "tube", "name": "Tube"},
            {"id": "jerigen", "name": "Jerigen"},
            {"id": "rollon", "name": "Roll On"},
            {"id": "makeup", "name": "Kemasan Make Up"},
            {"id": "aksesoris", "name": "Tutup & Pump"},
        ],
        "packaging": pack_pub,
    }
    (ROOT / "assets" / "js" / "catalog.js").write_text(
        "/* Dibuat otomatis oleh tools/build_data.py — jangan edit manual. Tidak berisi harga. */\n"
        "window.CATALOG = " + json.dumps(catalog, ensure_ascii=False, separators=(",", ":")) + ";\n")

    prices = {"version": stamp, "settings": settings, "bases": bases_sec,
              "actives": actives_sec, "packaging": pack_sec}
    (SRC / "prices.plain.json").write_text(json.dumps(prices, ensure_ascii=False, indent=1))
    (DATA / "prices.enc.json").write_text(json.dumps(encrypt(prices, password)))

    types = {}
    for p in pack_pub:
        types[p["type"]] = types.get(p["type"], 0) + 1
    noprice = [k for k, v in pack_sec.items() if not v["price"]]
    print(f"bases={len(bases_pub)} actives={len(actives_pub)} packaging={len(pack_pub)} types={types}")
    print("kemasan tanpa harga:", noprice)


if __name__ == "__main__":
    main()
